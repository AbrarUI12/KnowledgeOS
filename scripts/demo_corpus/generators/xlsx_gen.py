"""Generate XLSX documents using openpyxl for the demo corpus."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from scripts.demo_corpus.generators.determinism import (
    CREATOR,
    FIXED_DATE,
    normalize_ooxml,
)


def generate_xlsx(
    output_path: Path,
    title: str,
    sheets: list[dict[str, str | list[list[str]]]],
) -> None:
    """Generate an XLSX workbook with the given title and sheets.

    Each sheet is a dict with 'name' and 'rows' (list of lists of strings).

    Core properties are assigned before save; openpyxl overwrites ``modified``
    with the current time at save time, so ``normalize_ooxml`` afterwards rewrites
    the core-property timestamps and ZIP headers to fixed values, making the file
    byte-identical on every run (see ``determinism.py``).
    """
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    for sheet_spec in sheets:
        name = str(sheet_spec.get("name", "Sheet1"))
        rows = sheet_spec.get("rows", [])

        ws = wb.create_sheet(title=name)

        if isinstance(rows, list):
            for row_idx, row in enumerate(rows, 1):
                if isinstance(row, list):
                    for col_idx, cell_value in enumerate(row, 1):
                        cell = ws.cell(row=row_idx, column=col_idx, value=str(cell_value))
                        if row_idx == 1:
                            cell.font = Font(bold=True)
                elif isinstance(row, str):
                    ws.cell(row=row_idx, column=1, value=row)

    # Pin core properties. openpyxl honors `created`/`creator` but force-overwrites
    # `modified` at save time; normalize_ooxml() below pins that back to FIXED_DATE.
    wb.properties.creator = CREATOR
    wb.properties.lastModifiedBy = CREATOR
    wb.properties.created = FIXED_DATE
    wb.properties.modified = FIXED_DATE

    wb.save(str(output_path))
    normalize_ooxml(output_path)
