"""Generate DOCX documents using python-docx for the demo corpus."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from scripts.demo_corpus.generators.determinism import (
    CREATOR,
    FIXED_DATE,
    normalize_ooxml,
)


def generate_docx(
    output_path: Path,
    title: str,
    sections: list[dict[str, str | list[str]]],
) -> None:
    """Generate a DOCX with the given title and sections.

    Core properties are pinned and the ZIP container is normalized so the file is
    byte-identical on every run (see ``determinism.py``).
    """
    doc = Document()

    # Title
    doc.add_heading(title, level=0)

    for section in sections:
        heading = section.get("heading", "")
        content = section.get("content", [])

        if heading:
            doc.add_heading(str(heading), level=1)

        if isinstance(content, list):
            for para_text in content:
                doc.add_paragraph(str(para_text))
        elif isinstance(content, str):
            doc.add_paragraph(content)

    # Pin core properties (python-docx honors these on save).
    core = doc.core_properties
    core.author = CREATOR
    core.last_modified_by = CREATOR
    core.created = FIXED_DATE
    core.modified = FIXED_DATE

    doc.save(str(output_path))
    normalize_ooxml(output_path)
