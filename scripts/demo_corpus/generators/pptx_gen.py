"""Generate PPTX documents using python-pptx for the demo corpus."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation


def generate_pptx(
    output_path: Path,
    title: str,
    slides: list[dict[str, str | list[str]]],
) -> None:
    """Generate a PPTX presentation with the given title and slides.

    Each slide is a dict with 'title' and 'bullets' (list of strings).
    """
    prs = Presentation()

    # Title slide
    title_layout = prs.slide_layouts[0]  # Title Slide
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = "Demo Corpus"

    # Content slides
    bullet_layout = prs.slide_layouts[1]  # Title and Content
    for slide_spec in slides:
        slide_title = str(slide_spec.get("title", ""))
        bullets = slide_spec.get("bullets", [])

        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = slide_title

        if isinstance(bullets, list):
            body = slide.placeholders[1].text_frame
            body.clear()
            for bullet in bullets:
                p = body.add_paragraph()
                p.text = str(bullet)

    prs.save(str(output_path))
