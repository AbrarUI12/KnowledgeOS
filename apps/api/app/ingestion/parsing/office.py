"""DB-free Office/tabular adapters preserving source positions and reading order."""

import csv
import json
import tempfile
from collections.abc import Iterable, Iterator
from io import BytesIO, StringIO
from itertools import islice
from pathlib import Path
from typing import Any, cast

from charset_normalizer import from_bytes
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape

from app.ingestion.canonical import Block, BlockType, CanonicalDocument, TableData
from app.ingestion.parsing.base import ParserError
from app.ingestion.parsing.docling import block_id
from app.ingestion.parsing.isolation import run_isolated

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class OfficeParser:
    name = "office"
    version = "1"

    def __init__(
        self,
        *,
        max_rows: int = 50_000,
        max_columns: int = 200,
        row_group_size: int = 30,
        max_scanned_rows: int = 250_000,
        max_sheets: int = 100,
        timeout_seconds: float = 120,
        memory_bytes: int = 4 * 1024**3,
    ) -> None:
        if min(max_rows, max_columns, row_group_size, max_scanned_rows, max_sheets) < 1:
            raise ValueError("tabular limits must be positive")
        self.max_rows = max_rows
        self.max_columns = max_columns
        self.row_group_size = row_group_size
        self.max_scanned_rows = max_scanned_rows
        self.max_sheets = max_sheets
        self.timeout_seconds = timeout_seconds
        self.memory_bytes = memory_bytes

    def supports(self, mime_type: str) -> bool:
        return mime_type in {PPTX, XLSX, "text/csv", "text/plain"}

    def parse(self, data: bytes, *, filename: str, mime_type: str) -> CanonicalDocument:
        if not self.supports(mime_type):
            raise ParserError("Unsupported format.", code="unsupported_format")
        if not data or len(data) > 50 * 1024**2:
            raise ParserError("Invalid document size.", code="parser_input_limit")
        with tempfile.TemporaryDirectory(prefix="kos-office-") as temporary:
            directory = Path(temporary)
            source, output = directory / "source", directory / "result.json"
            source.write_bytes(data)
            request = directory / "request.json"
            request.write_text(
                json.dumps(
                    {
                        "source": str(source),
                        "output": str(output),
                        "filename": filename,
                        "mime_type": mime_type,
                        "max_rows": self.max_rows,
                        "max_columns": self.max_columns,
                        "row_group_size": self.row_group_size,
                        "max_scanned_rows": self.max_scanned_rows,
                        "max_sheets": self.max_sheets,
                        "memory_bytes": self.memory_bytes,
                    }
                ),
                encoding="utf-8",
            )
            run_isolated(
                "app.ingestion.parsing.office_worker",
                [str(request)],
                timeout_seconds=self.timeout_seconds,
                memory_bytes=self.memory_bytes,
            )
            if not output.exists() or output.stat().st_size > 64 * 1024**2:
                raise ParserError(
                    "Parser output is missing or too large.", code="parser_output_limit"
                )
            result = json.loads(output.read_bytes())
            if "error" in result:
                raise ParserError("Office document could not be parsed.", code=result["error"])
            return CanonicalDocument.model_validate(result["document"])

    def _parse_document(self, data: bytes, *, filename: str, mime_type: str) -> CanonicalDocument:
        if not self.supports(mime_type):
            raise ParserError("Unsupported format.", code="unsupported_format")
        if not data or len(data) > 50 * 1024**2:
            raise ParserError("Invalid document size.", code="parser_input_limit")
        document = CanonicalDocument(
            title=Path(filename).stem,
            mime_type=mime_type,
            parser_name=self.name,
            parser_version=self.version,
        )
        try:
            if mime_type == PPTX:
                self._slides(data, document)
            elif mime_type == XLSX:
                workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
                try:
                    if len(workbook.worksheets) > self.max_sheets:
                        raise ParserError("Too many sheets.", code="parser_table_limit")
                    remaining_rows, remaining_scans = self.max_rows, self.max_scanned_rows
                    for sheet in workbook.worksheets:
                        used, scanned = self._rows(
                            sheet.iter_rows(values_only=True),
                            sheet.title,
                            document,
                            row_budget=remaining_rows,
                            scan_budget=remaining_scans,
                        )
                        remaining_rows -= used
                        remaining_scans -= scanned
                finally:
                    workbook.close()
            else:
                result = from_bytes(data).best()
                if result is None:
                    raise ParserError(
                        "Text encoding could not be determined.", code="invalid_encoding"
                    )
                text = str(result)
                if mime_type == "text/csv":
                    try:
                        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
                    except csv.Error:
                        dialect = csv.excel
                    self._rows(csv.reader(StringIO(text), dialect), "CSV", document)
                else:
                    for paragraph in text.split("\n\n"):
                        if paragraph.strip():
                            self._append(document, BlockType.PARAGRAPH, paragraph.strip())
        except ParserError:
            raise
        except Exception as error:
            raise ParserError(
                "Office document could not be parsed.", code="parser_failed"
            ) from error
        return document

    @staticmethod
    def _append(document: CanonicalDocument, kind: BlockType, text: str, **location: Any) -> None:
        order = len(document.blocks)
        document.blocks.append(
            Block(
                block_id=block_id(order, text), type=kind, text=text, source_order=order, **location
            )
        )

    def _slides(self, data: bytes, document: CanonicalDocument) -> None:
        presentation = Presentation(BytesIO(data))
        if len(presentation.slides) > 500:
            raise ParserError("Too many slides.", code="parser_page_limit")
        for number, slide in enumerate(presentation.slides, 1):
            title = slide.shapes.title.text if slide.shapes.title else f"Slide {number}"
            self._append(document, BlockType.HEADING, title, slide=number, heading_level=1)
            for shape in self._leaf_shapes(slide.shapes):
                if shape == slide.shapes.title:
                    continue
                if shape.has_text_frame and cast(Shape, shape).text.strip():
                    self._append(
                        document,
                        BlockType.PARAGRAPH,
                        cast(Shape, shape).text.strip(),
                        slide=number,
                        heading_path=[title],
                    )
                if shape.has_table:
                    values = [
                        [cell.text for cell in row.cells]
                        for row in cast(GraphicFrame, shape).table.rows
                    ]
                    self._append(
                        document,
                        BlockType.TABLE,
                        "\n".join(" | ".join(row) for row in values),
                        slide=number,
                        heading_path=[title],
                        table=TableData(header=values[0], rows=values[1:]),
                    )
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes and notes.text.strip():
                    self._append(
                        document,
                        BlockType.PARAGRAPH,
                        notes.text.strip(),
                        slide=number,
                        heading_path=[title],
                        extra={"speaker_notes": True},
                    )
        document.metadata["slide_count"] = len(presentation.slides)

    @classmethod
    def _leaf_shapes(cls, shapes: Iterable[BaseShape], depth: int = 0) -> Iterator[BaseShape]:
        if depth > 64:
            raise ParserError("Shape nesting limit exceeded.", code="parser_shape_limit")
        for shape in sorted(shapes, key=lambda item: (item.top or 0, item.left or 0)):
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP and isinstance(shape, GroupShape):
                yield from cls._leaf_shapes(shape.shapes, depth + 1)
            else:
                yield shape

    def _rows(
        self,
        rows: Any,
        sheet: str,
        document: CanonicalDocument,
        *,
        row_budget: int | None = None,
        scan_budget: int | None = None,
    ) -> tuple[int, int]:
        row_budget = self.max_rows if row_budget is None else row_budget
        scan_budget = self.max_scanned_rows if scan_budget is None else scan_budget
        group: list[list[str]] = []
        header: list[str] = []
        start = end = used = scanned = 0
        for number, row in enumerate(islice(rows, scan_budget + 1), 1):
            scanned = number
            if number > scan_budget or len(row) > self.max_columns:
                raise ParserError("Tabular limits exceeded.", code="parser_table_limit")
            values = ["" if value is None else str(value) for value in row]
            if not any(value.strip() for value in values):
                continue
            used += 1
            if used > row_budget:
                raise ParserError("Tabular limits exceeded.", code="parser_table_limit")
            if used == 1 and any(isinstance(value, str) for value in row):
                header = values
                continue
            if not group:
                start = number
            end = number
            group.append(values)
            if len(group) == self.row_group_size:
                self._table_group(document, sheet, header, group, start, number)
                group = []
        if group:
            self._table_group(document, sheet, header, group, start, end)
        document.metadata["tabular_rag_not_analytics"] = True
        return used, scanned

    def _table_group(
        self,
        document: CanonicalDocument,
        sheet: str,
        header: list[str],
        rows: list[list[str]],
        start: int,
        end: int,
    ) -> None:
        text = "\n".join(" | ".join(row) for row in ([header] if header else []) + rows)
        self._append(
            document,
            BlockType.TABLE,
            text,
            sheet=sheet,
            row_start=start,
            row_end=end,
            table=TableData(header=list(header), rows=rows),
        )
