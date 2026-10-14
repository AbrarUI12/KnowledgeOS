from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from syrupy.assertion import SnapshotAssertion
from syrupy.types import PropertyName, PropertyPath

from app.ingestion.canonical import BlockType
from app.ingestion.parsing.base import ParserError
from app.ingestion.parsing.docling import DoclingParser
from app.ingestion.parsing.isolation import run_isolated

pytestmark = pytest.mark.unit
CORPUS = Path(__file__).resolve().parents[3] / "evals" / "fixtures" / "corpus" / "demo"
# Model/PDFium floating-point boxes vary by <0.5pt between Windows and Linux; all other
# canonical fields remain exact. Keep a fixed spatial reference rather than dropping bboxes.
PDF_BOXES = [
    (125.95, 83.44, 485.35, 102.18),
    (79.314, 152.896, 351.15, 169.762),
    (78.23, 190.72, 429.48, 224.09),
    (78.0, 238.72, 371.70, 248.09),
    (79.314, 259.112, 251.754, 275.762),
    (78.79, 296.84, 316.66, 306.09),
    (78.74, 320.72, 338.34, 330.09),
    (79.314, 340.896, 208.158, 354.198),
    (78.50, 378.72, 277.21, 387.98),
    (78.74, 402.81, 242.23, 411.98),
]


def exclude_bbox_coordinate(*, prop: PropertyName, path: PropertyPath) -> bool:
    return prop in {"x0", "y0", "x1", "y1"} and any(name == "bbox" for name, _ in path)


@pytest.mark.parametrize(
    ("filename", "mime_type"),
    [
        ("acme_hr_handbook_2026.pdf", "application/pdf"),
        (
            "acme_eng_runbook.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        ("acme_support_troubleshooting.html", "text/html"),
        ("acme_eng_architecture.md", "text/markdown"),
    ],
)
def test_real_format_canonical_snapshot(
    filename: str, mime_type: str, snapshot: SnapshotAssertion
) -> None:
    assets = Path(__file__).resolve().parents[1] / ".venv" / "docling-assets"
    document = DoclingParser(artifacts_path=assets if assets.exists() else None).parse(
        (CORPUS / filename).read_bytes(), filename=filename, mime_type=mime_type
    )
    assert document.blocks
    assert [b.source_order for b in document.blocks] == list(range(len(document.blocks)))
    if mime_type == "application/pdf":
        assert document.model_dump(mode="json") == snapshot(exclude=exclude_bbox_coordinate)
        for block, expected_box in zip(document.blocks, PDF_BOXES, strict=True):
            assert len(block.bbox) == 1
            box = block.bbox[0]
            assert box.page == 1
            assert (box.x0, box.y0, box.x1, box.y1) == pytest.approx(expected_box, rel=0, abs=0.5)
            assert 0 <= box.x0 < box.x1 <= 612
            assert 0 <= box.y0 < box.y1 <= 792
    else:
        assert document.model_dump(mode="json") == snapshot
    if mime_type == "application/pdf":
        assert document.page_count
        assert all(block.page for block in document.blocks)
        assert any(block.bbox for block in document.blocks)
        text = " ".join(block.text for block in document.blocks)
        for phrase in (
            "16 weeks of paid parental leave",
            "HR-142-B",
            "previous 12-week allowance",
            "completed 6 months of service",
            "4 days per week",
            "Manager approval is required",
        ):
            assert phrase in text
    if filename.endswith(".md"):
        assert any(block.type is BlockType.HEADING for block in document.blocks)


def test_real_child_timeout() -> None:
    with pytest.raises(ParserError) as error:
        run_isolated("tests.parser_probe", ["hang"], timeout_seconds=0.2, memory_bytes=1024**3)
    assert error.value.code == "parser_timeout"
    assert error.value.retryable


def test_docx_hierarchy_lists_and_structured_table() -> None:
    source = Document()
    source.add_heading("Policy", level=1)
    source.add_heading("Eligibility", level=2)
    source.add_paragraph("Active employees", style="List Bullet")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Category"
    table.cell(0, 1).text = "Weeks"
    table.cell(1, 0).text = "Employee"
    table.cell(1, 1).text = "16"
    output = BytesIO()
    source.save(output)
    document = DoclingParser().parse(
        output.getvalue(),
        filename="policy.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    listing = next(block for block in document.blocks if block.type is BlockType.LIST)
    assert listing.heading_path == ["Policy", "Eligibility"]
    parsed_table = next(block for block in document.blocks if block.type is BlockType.TABLE)
    assert parsed_table.table is not None
    assert ["Employee", "16"] in parsed_table.table.rows
    again = DoclingParser().parse(
        output.getvalue(),
        filename="policy.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert document == again


def test_real_child_memory_limit() -> None:
    with pytest.raises(ParserError) as error:
        run_isolated(
            "tests.parser_probe", ["allocate"], timeout_seconds=5, memory_bytes=32 * 1024**2
        )
    assert error.value.code == "parser_memory_limit"


def test_unsupported_and_empty_inputs_fail_before_launch() -> None:
    parser = DoclingParser()
    with pytest.raises(ParserError, match="Unsupported"):
        parser.parse(b"content", filename="input.exe", mime_type="application/octet-stream")
    with pytest.raises(ParserError, match="empty"):
        parser.parse(b"", filename="empty.pdf", mime_type="application/pdf")
