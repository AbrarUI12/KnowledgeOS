from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches

from app.ingestion.canonical import Block, BlockType, BoundingBox, CanonicalDocument
from app.ingestion.normalize import normalize_document
from app.ingestion.parsing.base import ParserError
from app.ingestion.parsing.office import PPTX, XLSX, OfficeParser

pytestmark = pytest.mark.unit
CORPUS = Path(__file__).resolve().parents[3] / "evals" / "fixtures" / "corpus" / "demo"


@pytest.mark.parametrize(
    ("filename", "mime"),
    [
        ("acme_eng_architecture_slides.pptx", PPTX),
        ("acme_eng_team_metrics.xlsx", XLSX),
        ("acme_eng_metrics.csv", "text/csv"),
    ],
)
def test_real_office_snapshot(filename: str, mime: str, snapshot: object) -> None:
    result = OfficeParser(row_group_size=2).parse(
        (CORPUS / filename).read_bytes(), filename=filename, mime_type=mime
    )
    assert result.blocks
    assert result.model_dump(mode="json") == snapshot
    assert (
        OfficeParser(row_group_size=2).parse(
            (CORPUS / filename).read_bytes(), filename=filename, mime_type=mime
        )
        == result
    )


def test_semicolon_csv_row_ranges_and_caps() -> None:
    result = OfficeParser(row_group_size=1).parse(
        b"Name;Weeks\nAlice;16\nBob;12", filename="leave.csv", mime_type="text/csv"
    )
    assert result.blocks[0].row_start == 2
    assert result.blocks[1].row_end == 3
    assert result.blocks[0].table is not None
    assert result.blocks[0].table.header == ["Name", "Weeks"]
    with pytest.raises(ParserError) as error:
        OfficeParser(max_columns=1).parse(b"a,b\n1,2", filename="wide.csv", mime_type="text/csv")
    assert error.value.code == "parser_table_limit"


def test_office_conversion_runs_in_a_bounded_child() -> None:
    with pytest.raises(ParserError) as error:
        OfficeParser(timeout_seconds=0.01).parse(
            b"Name;Weeks\nAlice;16", filename="leave.csv", mime_type="text/csv"
        )
    assert error.value.code == "parser_timeout"
    assert error.value.retryable


def test_sparse_workbook_skips_empty_rows_preserving_physical_ranges() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Name", "Weeks"])
    sheet.cell(10, 1, "Alice")
    sheet.cell(10, 2, 16)
    sheet.cell(50, 1, "Bob")
    sheet.cell(50, 2, 12)
    sheet.cell(60_000, 2).font = Font(bold=True)
    data = BytesIO()
    workbook.save(data)
    result = OfficeParser(max_rows=3, row_group_size=2).parse(
        data.getvalue(), filename="sparse.xlsx", mime_type=XLSX
    )
    assert len(result.blocks) == 1
    block = result.blocks[0]
    assert (block.row_start, block.row_end) == (10, 50)
    assert block.table is not None
    assert block.table.rows == [["Alice", "16"], ["Bob", "12"]]
    assert block.table.header == ["Name", "Weeks"]


def test_workbook_row_limit_is_global_across_sheets() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Name"])
    sheet.append(["Alice"])
    workbook.create_sheet("second").append(["Bob"])
    data = BytesIO()
    workbook.save(data)
    with pytest.raises(ParserError) as error:
        OfficeParser(max_rows=2).parse(data.getvalue(), filename="many.xlsx", mime_type=XLSX)
    assert error.value.code == "parser_table_limit"
    assert not error.value.retryable


@pytest.mark.parametrize("limit", ["sheets", "scans"])
def test_workbook_sheet_and_physical_scan_caps_are_permanent(limit: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.cell(2, 1, "First")
    workbook.create_sheet("second").cell(2, 1, "Second")
    data = BytesIO()
    workbook.save(data)
    parser = OfficeParser(max_sheets=1) if limit == "sheets" else OfficeParser(max_scanned_rows=3)
    with pytest.raises(ParserError) as error:
        parser.parse(data.getvalue(), filename="caps.xlsx", mime_type=XLSX)
    assert error.value.code == "parser_table_limit"
    assert not error.value.retryable


def test_sparse_csv_header_and_last_partial_group_have_physical_ranges() -> None:
    result = OfficeParser(row_group_size=2).parse(
        b"\nName,Weeks\n\nAlice,16\n\nBob,12\n\nCarol,10\n",
        filename="sparse.csv",
        mime_type="text/csv",
    )
    assert [(block.row_start, block.row_end) for block in result.blocks] == [(4, 6), (8, 8)]
    assert all(block.table and block.table.header == ["Name", "Weeks"] for block in result.blocks)


def test_grouped_slide_text_and_table_are_not_lost() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    nested = group.shapes.add_group_shape()
    text = nested.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    text.text = "Grouped engineering label"
    table_shape = slide.shapes.add_table(2, 1, Inches(1), Inches(3), Inches(2), Inches(1))
    group.shapes.add_group_shape([table_shape])
    table = table_shape.table
    table.cell(0, 0).text = "Team"
    table.cell(1, 0).text = "Engineering"
    data = BytesIO()
    presentation.save(data)
    result = OfficeParser().parse(data.getvalue(), filename="groups.pptx", mime_type=PPTX)
    assert [block.text for block in result.blocks] == [
        "Slide 1",
        "Grouped engineering label",
        "Team\nEngineering",
    ]
    assert all(block.slide == 1 for block in result.blocks)
    assert result.blocks[-1].table is not None
    assert result.blocks[-1].table.rows == [["Engineering"]]


def test_slide_shape_without_position_still_parses() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    shape.text = "Unpositioned label"
    # Real malformed-but-readable OOXML shape: xfrm is optional, not a mocked parser.
    properties = shape._element.spPr
    assert properties.xfrm is not None
    properties.remove(properties.xfrm)
    assert shape.top is None
    data = BytesIO()
    presentation.save(data)
    result = OfficeParser().parse(data.getvalue(), filename="position.pptx", mime_type=PPTX)
    assert result.blocks[1].text == "Unpositioned label"


def test_normalization_preserves_code_and_original_input() -> None:
    document = CanonicalDocument(
        title="test",
        mime_type="text/plain",
        parser_name="fixture",
        parser_version="1",
        blocks=[
            Block(
                block_id="original1",
                source_order=0,
                type=BlockType.PARAGRAPH,
                text="  cafe\u0301\u202e\u200b   policy  ",
            ),
            Block(
                block_id="original2",
                source_order=1,
                type=BlockType.CODE,
                text="if x:\n    return x",
            ),
        ],
    )
    normalized = normalize_document(document, retain_raw=True)
    assert normalized.blocks[0].text == "café policy"
    assert normalized.blocks[0].raw_text == document.blocks[0].text
    assert normalized.blocks[1].text == "if x:\n    return x"
    assert document.normalization_version is None
    assert normalize_document(normalized, retain_raw=True) == normalized


def test_repeated_top_margin_is_furniture_not_body_clauses() -> None:
    blocks = [
        Block(
            block_id=f"header00{page}",
            source_order=page - 1,
            page=page,
            type=BlockType.PARAGRAPH,
            text="Acme Handbook",
            bbox=[BoundingBox(page=page, x0=10, y0=10, x1=100, y1=20)],
        )
        for page in (1, 2)
    ]
    blocks.append(
        Block(
            block_id="body0001",
            source_order=2,
            page=2,
            type=BlockType.PARAGRAPH,
            text="Acme Handbook",
            bbox=[BoundingBox(page=2, x0=10, y0=100, x1=100, y1=110)],
        )
    )
    document = CanonicalDocument(
        title="test",
        mime_type="application/pdf",
        parser_name="fixture",
        parser_version="1",
        blocks=blocks,
    )
    result = normalize_document(document)
    assert all(block.type is BlockType.FURNITURE for block in result.blocks[:2])
    assert result.blocks[2].type is BlockType.PARAGRAPH
